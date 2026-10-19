"""Check source fidelity, native direction metadata and PDF font/extraction evidence.
Uses stdlib for DOCX/PPTX, openpyxl only for XLSX and pypdf only for exported PDFs.
A failed PDF round trip is
reported independently from source/structure checks, not hidden as a pass.
"""
import argparse,hashlib,importlib.metadata,json,re,zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

def main():
    ap=argparse.ArgumentParser();ap.add_argument('--input',type=Path,required=True);args=ap.parse_args();root=args.input.resolve()
    manifest=json.loads((root/'fixtures.json').read_text(encoding='utf-8'));cases=manifest['cases'];checks=[]
    def check(name,ok):checks.append({'name':name,'pass':bool(ok)})
    selected=set(manifest.get('formats') or [Path(f).suffix[1:] for f in manifest['sha256']])
    if selected=={'manifest'}:ap.error('Manifest-only fixtures use the WPF renderer; no source format selected')
    packages=set()
    for file,digest in manifest['sha256'].items():check('fixture hash '+file,hashlib.sha256((root/file).read_bytes()).hexdigest()==digest)
    ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main','a':'http://schemas.openxmlformats.org/drawingml/2006/main'}
    if 'docx' in selected:
        with zipfile.ZipFile(root/'fixture.docx') as z:
            doc=ET.fromstring(z.read('word/document.xml'));ps=doc.findall('.//w:body/w:p',ns)
            text=[''.join(n.text or '' for n in p.findall('.//w:t',ns)) for p in ps]
            check('DOCX logical paragraph content',text[1:]==[c['logical'] for c in cases])
            check('DOCX bidi on every paragraph',all(p.find('w:pPr/w:bidi',ns) is not None for p in ps))
            check('DOCX logical start alignment',all(p.find('w:pPr/w:jc',ns).get('{'+ns['w']+'}val')=='start' for p in ps))
    if 'pptx' in selected:
        with zipfile.ZipFile(root/'fixture.pptx') as z:
            slides=sorted(n for n in z.namelist() if re.fullmatch(r'ppt/slides/slide\d+\.xml',n));text=[];paragraphs=[];shapes=0
            for name in slides:
                slide=ET.fromstring(z.read(name));ps=slide.findall('.//a:p',ns);paragraphs+=ps
                text.extend(''.join(n.text or '' for n in p.findall('.//a:t',ns)) for p in ps)
                shapes+=len(slide.findall('.//{http://schemas.openxmlformats.org/presentationml/2006/main}sp'))
            check('PPTX logical paragraph content',text==[c['logical'] for c in cases])
            check('PPTX RTL paragraph metadata',all(p.find('a:pPr',ns).get('rtl')=='1' for p in paragraphs))
            check('PPTX native paragraphs not word boxes',len(slides)==2 and shapes==2)
    if 'xlsx' in selected:
        from openpyxl import load_workbook
        wb=load_workbook(root/'fixture.xlsx');ws=wb.worksheets[0]
        check('XLSX RTL sheet view',ws.sheet_view.rightToLeft)
        check('XLSX clean canonical source',[wb['canonical'].cell(i+2,2).value for i in range(len(cases))]==[c['logical'] for c in cases])
        check('XLSX display maps to canonical',[ws.cell(i+2,2).value.replace('\u200e','') for i in range(len(cases))]==[c['logical'] for c in cases])
        check('XLSX numeric values remain numeric',all(ws.cell(i+2,3).value==-12.5 and ws.cell(i+2,3).data_type=='n' for i in range(len(cases))))
        check('XLSX RTL/LTR cells scoped',all(ws.cell(i+2,2).alignment.readingOrder==2 and ws.cell(i+2,4).alignment.readingOrder==1 for i in range(len(cases))))
        check('XLSX no unsupported isolate controls',all(not any(ch in ws.cell(i+2,2).value for ch in '\u2066\u2067\u2068\u2069') for i in range(len(cases))))
        packages.add('openpyxl')
    pdf_names={'html':'browser.pdf','docx':'word.pdf','pptx':'powerpoint.pdf','xlsx':'excel.pdf'}
    pdf_files=[root/pdf_names[f] for f in sorted(selected) if f in pdf_names and (root/pdf_names[f]).exists()]
    if pdf_files:
        from pypdf import PdfReader
        packages.add('pypdf')
    pdf=[]
    def compact(s):return re.sub(r'\s+','',s)
    for file in pdf_files:
        reader=PdfReader(file);text='\n'.join(p.extract_text() or '' for p in reader.pages);fonts=[]
        for page in reader.pages:
            for ref in page.get('/Resources',{}).get('/Font',{}).values():
                f=ref.get_object();desc=f.get('/FontDescriptor')
                if desc is None and f.get('/DescendantFonts'):desc=f['/DescendantFonts'][0].get_object().get('/FontDescriptor')
                desc=desc.get_object() if desc else {}
                fonts.append({'name':str(f.get('/BaseFont')),'embedded':any(k in desc for k in ['/FontFile','/FontFile2','/FontFile3']),'toUnicode':bool(f.get('/ToUnicode'))})
        roundtrip=[{'case':c['id'],'exactAfterWhitespaceNormalization':compact(c['logical']) in compact(text)} for c in cases]
        check(file.name+' embedded fonts',bool(fonts) and all(f['embedded'] for f in fonts))
        pdf.append({'file':file.name,'pages':len(reader.pages),'sha256':hashlib.sha256(file.read_bytes()).hexdigest(),'fonts':fonts,'extractor':'pypdf '+importlib.metadata.version('pypdf'),'roundtrip':roundtrip,'viewerFind':'not tested','copyPaste':'not tested'})
    report={'checks':checks,'pdf':pdf,'selectedFormats':sorted(selected),'packages':{n:importlib.metadata.version(n) for n in sorted(packages)}}
    (root/'artifact-results.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
    failures=[x['name'] for x in checks if not x['pass']]
    print(json.dumps({'checks':len(checks),'passed':len(checks)-len(failures),'failed':failures,'pdfExactRoundtrips':{x['file']:str(sum(r['exactAfterWhitespaceNormalization'] for r in x['roundtrip']))+'/'+str(len(cases)) for x in pdf}}))
    raise SystemExit(bool(failures))

if __name__=='__main__':main()
