"""Create synthetic RTL fixtures in logical Unicode order; never edit source artifacts.
Usage: python build_fixtures.py --out DIRECTORY --formats html [--font Arial]
Only selected formats are generated/imported. Default: html. Use manifest for WPF.
"""
import argparse
import hashlib
import html
import json
import platform
import re
from pathlib import Path

CASES = [
    ('hebrew', [('בדיקת עברית תקינה מראש.', 'rtl')]),
    ('mixed', [('המערכת ', 'rtl'), ('Revit 2025', 'ltr'), (' מציגה טקסט קריא.', 'rtl')]),
    ('latin-start', [('PRJ07', 'ltr'), (' — תוכנית עבודה בעברית', 'rtl')]),
    ('code', [('קוד הרכיב הוא ', 'rtl'), ('ABC-123/07', 'ltr'), ('; יש לשמור את המקף והלוכסן.', 'rtl')]),
    ('number', [('הפרש המדידה: ', 'rtl'), ('-12.50 mm', 'ltr'), ('.', 'rtl')]),
    ('date', [('מועד הבדיקה: ', 'rtl'), ('19.10.2026 14:30', 'ltr'), ('.', 'rtl')]),
    ('brackets', [('הבדיקה הושלמה ', 'rtl'), ('(QA-17)', 'ltr'), (' בהצלחה.', 'rtl')]),
    ('punctuation', [('מה נבדק? עברית, מספרים ואנגלית; נשמרו: פסיק, נקודתיים ונקודה.', 'rtl')]),
    ('units', [('נפח הבטון הוא ', 'rtl'), ('1,234.50 m³', 'ltr'), (' ושיעור ההשלמה ', 'rtl'), ('95.2%', 'ltr'), ('.', 'rtl')]),
    ('path', [('קובץ העבודה: ', 'rtl'), ('C:\\Models\\PRJ07\\model.rvt', 'ltr'), ('.', 'rtl')]),
    ('url', [('כתובת לדוגמה: ', 'rtl'), ('https://example.com/a?x=12&y=34', 'ltr'), ('.', 'rtl')]),
    ('long', [('זוהי פסקה ארוכה לבדיקת שבירת שורות בעברית עם ', 'rtl'), ('API v2.1', 'ltr'), (' ושמירה על סדר קריא גם במסך צר. הסוגריים ', 'rtl'), ('(ABC-123)', 'ltr'), (' נשארים צמודים למזהה והפיסוק נשמר בסוף המשפט.', 'rtl')]),
    ('niqqud', [('שָׁלוֹם, עולם — בדיקת ניקוד ואותיות סופיות ךםןףץ.', 'rtl')]),
]

# Web-only line-break cases: each renders its good form; the bad form waits in a
# <template data-line-bad> for scratch negative controls (see reproducible-checks.md).
LINE_CASES = [
    ('line-status',
     '<p class="status-row"><span class="status-icon" aria-hidden="true"></span><span>נשמרו 8\u00a0קבצים במכשיר, אפשר לפתוח אותם גם בלי חיבור\u00a0·\u00a012\u00a0מ״ב</span></p>',
     '<p class="status-row"><span class="status-icon" aria-hidden="true"></span>נשמרו 8 קבצים במכשיר, אפשר לפתוח אותם גם בלי חיבור · <bdi>12 מ״ב</bdi></p>'),
    ('line-count',
     '<div class="cat-row"><button type="button" class="cat-name"><span>מאפים, עוגות וקינוחים לשבת ולחגים</span><small class="cat-count">4\u00a0מתכונים</small></button></div>',
     '<div class="cat-row"><button type="button" class="cat-name">מאפים, עוגות וקינוחים לשבת ולחגים</button> · 4</div>'),
    ('line-balance',
     '<div class="error-col"><p class="error-text" style="max-inline-size:100%">לא הצלחנו לשלוח קוד אל <bdi dir="ltr">grandmother.recipes.family@example.com</bdi>. נסו שוב.</p></div>',
     '<div class="error-col"><p class="error-text">לא הצלחנו לשלוח קוד אל <bdi dir="ltr">grandmother.recipes.family@example.com</bdi>. נסו שוב.</p></div>'),
]

def plain(runs):
    return ''.join(t for t, _ in runs)

def write_html(out, font):
    rows=[]
    for cid,runs in CASES:
        # A short quantity ("-12.50 mm") stays on one line without changing the shared canonical text.
        body=''.join('<bdi dir="ltr" data-ltr'+(' class="value"' if re.fullmatch(r'[-+]?[\d.,]+ \S{1,4}',t) else '')+'>'+html.escape(t)+'</bdi>' if d=='ltr' else '<span>'+html.escape(t)+'</span>' for t,d in runs)
        rows.append(f'<section class="case" id="{cid}"><small dir="ltr">{cid}</small><p dir="rtl">{body}</p></section>')
    lines=''.join(f'<section class="line-case" id="{cid}"><small dir="ltr">{cid}</small><div data-line-good>{good}</div><template data-line-bad>{bad}</template></section>' for cid,good,bad in LINE_CASES)
    page='''<!doctype html><html lang="he" dir="rtl"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>בדיקות עברית וכיווניות</title>
<style>
*{box-sizing:border-box}body{font-family:FONT,sans-serif;margin:0;color:#162d40;background:#f2f5f8;font-size:1.1875rem;line-height:1.65}main{max-inline-size:1050px;margin-inline:auto;padding:24px;background:white}h1{font-size:30px;line-height:1.3}h2{font-size:24px}.case{border-block-end:1px solid #ddd;padding-block:5px}.case p{margin-block:4px 10px;text-align:start;overflow-wrap:anywhere}.case small{display:block;color:#657383;font-size:12px;text-align:end}bdi{unicode-bidi:isolate}.value{white-space:nowrap}table{border-collapse:collapse;min-inline-size:600px;inline-size:100%}td,th{border:1px solid #ccd4dc;padding:8px;text-align:start}th{background:#eaf3f7}.table-scroll{overflow:auto;max-inline-size:100%}nav{display:grid;direction:ltr;grid-template-columns:1fr auto 1fr;gap:8px;margin-block:18px}#prev{grid-column:3;grid-row:1}#next{grid-column:1;grid-row:1}#counter{grid-column:2;grid-row:1;align-self:center}button{font:inherit;min-block-size:44px;border:1px solid #7496b0;background:#edf4fa;border-radius:6px}code,pre{direction:ltr;unicode-bidi:isolate;text-align:left;display:block;white-space:pre-wrap;overflow-wrap:anywhere}svg{max-inline-size:100%;block-size:auto}.negative{direction:ltr;text-align:right;background:#fff0ef}.line-case{border-block-end:1px solid #ddd;padding-block:5px}.line-case small{display:block;color:#657383;font-size:12px;text-align:end}.status-row{display:flex;gap:.5rem;align-items:flex-start;margin-block:4px 10px}.status-icon{flex:none;inline-size:1.25em;block-size:1.25em;margin-block-start:.2em;border-radius:50%;background:#7496b0}.cat-row{display:flex;align-items:center;gap:.4rem;max-inline-size:13em;margin-block:4px 10px}.cat-name{text-align:start;padding:4px 10px}.cat-count{display:block;color:#566475;font-size:.85em}.error-col{display:flex;flex-direction:column;align-items:center;text-align:center}.error-text{text-wrap:balance;overflow-wrap:break-word;margin-block:4px 10px}details{margin-block:18px}@media(max-width:500px){main{padding:16px}h1{font-size:25px}body{font-size:1.125rem}}@page{size:A4;margin:12mm}@media print{body{background:white;font-size:11pt;line-height:1.4}main{padding:0;max-inline-size:none}h1{font-size:22pt}.case{break-inside:avoid}.case p{margin-block:2px 5px}.case small{font-size:8pt}nav,details,.screen-only{display:none}.table-scroll{overflow:visible}table{min-inline-size:0;font-size:10pt}h2{font-size:15pt}}
</style><main><h1 dir="rtl">כתיבה בעברית עם <bdi dir="ltr">RTL</bdi> תקין מראש</h1><p>המחרוזות נשמרות בסדר לוגי. אנגלית, מספרים וסוגריים מבודדים בשכבת התצוגה.</p>ROWS
<h2>טבלה עם יחידות וכותרות</h2><div class="table-scroll" dir="rtl"><table><thead><tr><th>תיאור</th><th>קוד <bdi dir="ltr">ID</bdi></th><th>נפח <bdi dir="ltr">(m³)</bdi></th><th>הפרש <bdi dir="ltr">(mm)</bdi></th></tr></thead><tbody><tr><td>קיר בטון</td><td dir="ltr">ABC-123</td><td dir="ltr">1,234.50</td><td dir="ltr">-12.50</td></tr></tbody></table></div>
<nav aria-label="ניווט לבדיקה"><button id="prev" dir="rtl"><span aria-hidden="true">→</span> הקודם</button><span id="counter" dir="ltr">1 / 3</span><button id="next" dir="rtl">הבא <span aria-hidden="true">←</span></button></nav>
<div class="screen-only" id="lines"><h2>שבירת שורות</h2>LINES</div>
<div class="screen-only"><h2>טקסט טכני וקוד</h2><pre dir="ltr">result = model.get_element("ABC-123");</pre><svg viewBox="0 0 350 100" width="350" height="100" role="img" aria-label="כותרת עברית ב SVG"><text x="330" y="38" direction="rtl" text-anchor="start" font-family="FONT" font-size="24">כותרת עברית בתרשים</text><text x="330" y="76" direction="ltr" text-anchor="end" font-family="FONT" font-size="24">API v2.1</text></svg></div>
<details open class="screen-only"><summary>בקרת כשל מכוונת</summary><p id="negative" class="negative" dir="ltr">PRJ07 — תוכנית עבודה בעברית (QA-17).</p><p>שורה זו מיושרת לימין אך כיוון הבסיס שלה שגוי. עליה להיכשל בבדיקת הכיווניות.</p></details></main>
<script>let n=1;const paint=()=>{document.querySelector('#counter').textContent=n+' / 3';document.querySelector('#prev').disabled=n===1;document.querySelector('#next').disabled=n===3};document.querySelector('#prev').onclick=()=>{n=Math.max(1,n-1);paint()};document.querySelector('#next').onclick=()=>{n=Math.min(3,n+1);paint()};document.addEventListener('keydown',e=>{if(e.key==='ArrowLeft'){n=Math.min(3,n+1);paint()}if(e.key==='ArrowRight'){n=Math.max(1,n-1);paint()}});paint();</script></html>'''
    page=page.replace('ROWS','\n'.join(rows)).replace('LINES',lines).replace('FONT',html.escape(font))
    page=page.replace('</style>','pre{white-space:pre;overflow:auto;overflow-wrap:normal}</style>')
    (out/'fixture.html').write_text(page,encoding='utf-8')

def write_docx(out,font,alignment='start'):
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    doc=Document()
    normal=doc.styles['Normal'];normal.font.name=font;normal.font.size=Pt(13)
    doc.styles['Title'].font.color.rgb=RGBColor(0,0,0)
    for st in [doc.styles['Normal'],doc.styles['Title']]:
        pr=st.element.find(qn('w:pPr'))
        if pr is not None:
            for border in list(pr.findall(qn('w:pBdr'))):pr.remove(border)
    def para(runs,style=None):
        p=doc.add_paragraph(style=style);p.alignment=WD_ALIGN_PARAGRAPH.LEFT
        bidi=OxmlElement('w:bidi');p._p.get_or_add_pPr().append(bidi)
        p._p.get_or_add_pPr().find(qn('w:jc')).set(qn('w:val'),alignment)
        for t,d in runs:
            r=p.add_run(t);r.font.name=font;r.font.rtl=(d=='rtl');r.font.color.rgb=RGBColor(0,0,0)
            pr=r._r.get_or_add_rPr();fonts=pr.find(qn('w:rFonts'))
            if fonts is None:fonts=OxmlElement('w:rFonts');pr.insert(0,fonts)
            for a in ['ascii','hAnsi','cs']:fonts.set(qn('w:'+a),font)
            lang=OxmlElement('w:lang');lang.set(qn('w:val'),'he-IL' if d=='rtl' else 'en-US');lang.set(qn('w:bidi'),'he-IL');pr.append(lang)
        return p
    para([('בדיקות עברית וכיווניות','rtl')],'Title')
    for _,runs in CASES:para(runs)
    doc.save(out/'fixture.docx')

def write_pptx(out,font):
    from pptx import Presentation
    from pptx.util import Inches,Pt
    from pptx.enum.text import PP_ALIGN
    from pptx.oxml.xmlchemy import OxmlElement
    prs=Presentation();prs.slide_width=Inches(13.333);prs.slide_height=Inches(7.5)
    for offset in [0,7]:
        slide=prs.slides.add_slide(prs.slide_layouts[6])
        box=slide.shapes.add_textbox(Inches(.6),Inches(.3),Inches(12.1),Inches(6.8));tf=box.text_frame;tf.word_wrap=True
        for idx,(_,runs) in enumerate(CASES[offset:offset+7]):
            p=tf.paragraphs[0] if idx==0 else tf.add_paragraph();p.alignment=PP_ALIGN.RIGHT
            pr=p._p.get_or_add_pPr();pr.set('rtl','1');pr.set('algn','r')
            p.space_after=Pt(13)
            for t,d in runs:
                r=p.add_run();r.text=t;r.font.name=font;r.font.size=Pt(21)
                rp=r._r.get_or_add_rPr();rp.set('lang','he-IL' if d=='rtl' else 'en-US')
                cs=OxmlElement('a:cs');cs.set('typeface',font);rp.append(cs)
                if d=='rtl':
                    rt=OxmlElement('a:rtl');rt.set('val','1');rp.append(rt)
    prs.save(out/'fixture.pptx')

def write_xlsx(out,font,display_mode='lrm'):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment,Font,PatternFill
    wb=Workbook();ws=wb.active;ws.title='בדיקות RTL';ws.sheet_view.rightToLeft=True
    canonical=wb.create_sheet('canonical');canonical.append(['case','logical text'])
    ws.append(['מקרה','טקסט בסדר לוגי','ערך מספרי','קוד טכני'])
    for i,(cid,runs) in enumerate(CASES,2):
        # This Excel 16 export displays LRI/PDI as visible boxes. Use tested
        # local LRM boundaries only in the presentation column; source stays clean.
        left,right={'lrm':('\u200e','\u200e'),'isolate':('\u2066','\u2069'),'plain':('','')}[display_mode]
        display=''.join(left+t+right if d=='ltr' else t for t,d in runs)
        ws.append([cid,display,-12.5,'ABC-123/07'])
        canonical.append([cid,plain(runs)])
        ws.cell(i,3).number_format='0.00'
        for col in range(1,5):ws.cell(i,col).alignment=Alignment(horizontal='right' if col==2 else 'left',vertical='center',wrap_text=True,readingOrder=2 if col==2 else 1)
        ws.row_dimensions[i].height=42
    for row in ws:
        for c in row:c.font=Font(name=font,size=13)
    for c in ws[1]:c.fill=PatternFill('solid',fgColor='EAF3F7');c.alignment=Alignment(horizontal='right',readingOrder=2)
    for col,width in [('A',17),('B',95),('C',18),('D',22)]:ws.column_dimensions[col].width=width
    ws.freeze_panes='A2';ws.sheet_properties.pageSetUpPr.fitToPage=True
    ws.page_setup.orientation='landscape';ws.page_setup.paperSize=ws.PAPERSIZE_A4;ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=1
    ws.print_options.horizontalCentered=True
    canonical.sheet_state='hidden'
    wb.save(out/'fixture.xlsx')

def main():
    ap=argparse.ArgumentParser();ap.add_argument('--out',type=Path,required=True);ap.add_argument('--font',default='Arial')
    ap.add_argument('--formats',nargs='+',choices=['html','docx','pptx','xlsx','manifest'],default=['html'])
    ap.add_argument('--word-alignment',choices=['start','left','right','end','center'],default='start')
    ap.add_argument('--excel-display',choices=['lrm','isolate','plain'],default='lrm');args=ap.parse_args()
    if 'manifest' in args.formats and len(args.formats)>1:ap.error('manifest is a standalone mode; every other format includes a manifest already')
    selected=list(dict.fromkeys(args.formats))
    out=args.out.resolve();out.mkdir(parents=True,exist_ok=True)
    if 'html' in selected:write_html(out,args.font)
    if 'docx' in selected:write_docx(out,args.font,args.word_alignment)
    if 'pptx' in selected:write_pptx(out,args.font)
    if 'xlsx' in selected:write_xlsx(out,args.font,args.excel_display)
    manifest={'python':platform.python_version(),'formats':selected,'font_requested':args.font,'word_alignment':args.word_alignment,'excel_display':args.excel_display,'cases':[{'id':cid,'logical':plain(r),'runs':r} for cid,r in CASES],'line_cases':[cid for cid,_,_ in LINE_CASES] if 'html' in selected else [],'sha256':{f'fixture.{fmt}':hashlib.sha256((out/f'fixture.{fmt}').read_bytes()).hexdigest() for fmt in selected if fmt!='manifest'}}
    (out/'fixtures.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf-8')
    print('Created '+', '.join(selected)+' fixture in '+str(out))

if __name__=='__main__':main()
